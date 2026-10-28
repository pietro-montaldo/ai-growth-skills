"""Build a ranked Excel workbook from a leads.json produced by the icp-lead-ranker skill.

Needs openpyxl (pip install openpyxl).

Usage:
    python3 build_xlsx.py <leads.json> <output_dir>

leads.json schema (see SKILL.md):
{
  "meta": {"source_file": "", "product": "", "project_slug": "2026-06-25-icp-lead-ranking"},
  "leads": [
    {"name": "", "company": "", "role": "", "location": "",
     "linkedin_url": "", "email": "", "email_status": "", "match_note": "",
     "fit": 0, "band": "Strong|Moderate|Weak",
     "top_signal": "", "rationale": "", "next_step": "",
     "research": "compact dossier text", "sources": ["url1", "url2"]}
  ]
}

Output:
    <project_slug>-leads-ranked.xlsx in <output_dir>
    (falls back to today's date if project_slug is missing)
"""
import json
import os
import re
import sys
from datetime import date as _date
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


# ---------------------- Styling ----------------------
FONT_NAME = "Arial"
HEADER_FILL = PatternFill("solid", start_color="1F3A5F")
HEADER_FONT = Font(name=FONT_NAME, bold=True, color="FFFFFF", size=11)
SECTION_FILL = PatternFill("solid", start_color="E8EEF7")
SECTION_FONT = Font(name=FONT_NAME, bold=True, size=11, color="1F3A5F")
BODY_FONT = Font(name=FONT_NAME, size=10)
BAND_FILLS = {
    "Strong": PatternFill("solid", start_color="C6EFCE"),
    "Moderate": PatternFill("solid", start_color="FFEB9C"),
    "Weak": PatternFill("solid", start_color="FFC7CE"),
}
THIN = Side(border_style="thin", color="CCCCCC")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center")


# ---------------------- Helpers ----------------------
def band_for(fit, given=None):
    if given in BAND_FILLS:
        return given
    try:
        f = int(fit)
    except (TypeError, ValueError):
        return "Weak"
    if f >= 4:
        return "Strong"
    if f == 3:
        return "Moderate"
    return "Weak"


def sanitize_tab(name, used):
    base = re.sub(r"[:\\/\?\*\[\]]", "", str(name) or "lead")[:28].strip() or "lead"
    tab = base
    i = 2
    while tab.lower() in used:
        suffix = f" {i}"
        tab = base[:28 - len(suffix)] + suffix
        i += 1
    used.add(tab.lower())
    return tab


def row_height(text, per_chars=6, lo=20, hi=220):
    return max(lo, min(hi, 16 + len(str(text or "")) // per_chars))


# ---------------------- Main ----------------------
def main():
    if len(sys.argv) != 3:
        print("Usage: python3 build_xlsx.py <leads.json> <output_dir>", file=sys.stderr)
        sys.exit(1)

    leads_path = os.path.abspath(sys.argv[1])
    out_dir = os.path.abspath(sys.argv[2])
    os.makedirs(out_dir, exist_ok=True)

    with open(leads_path) as f:
        data = json.load(f)

    meta = data.get("meta", {})
    leads = data.get("leads", [])
    if not leads:
        print("leads.json has no leads.", file=sys.stderr)
        sys.exit(1)

    # Sort by fit desc, then email-having first (tie-break), then name.
    def sort_key(l):
        try:
            fit = int(l.get("fit", 0))
        except (TypeError, ValueError):
            fit = 0
        has_email = 1 if (l.get("email") or "").strip() else 0
        return (-fit, -has_email, str(l.get("name", "")).lower())

    leads_sorted = sorted(leads, key=sort_key)

    wb = Workbook()

    # ---------------- Overview ----------------
    ws = wb.active
    ws.title = "Overview"
    headers = ["#", "Name", "Company", "Role", "Location", "Fit", "Band",
               "Top signal", "Rationale", "Email", "Next step", "LinkedIn"]
    ws.append(headers)
    for col_idx, _ in enumerate(headers, 1):
        c = ws.cell(row=1, column=col_idx)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.alignment = CENTER
        c.border = BORDER

    widths = [4, 22, 24, 26, 20, 6, 11, 44, 50, 26, 50, 32]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w

    for i, lead in enumerate(leads_sorted, start=1):
        row = i + 1
        band = band_for(lead.get("fit"), lead.get("band"))
        ws.cell(row=row, column=1, value=i).alignment = CENTER
        ws.cell(row=row, column=2, value=lead.get("name", ""))
        ws.cell(row=row, column=3, value=lead.get("company", ""))
        ws.cell(row=row, column=4, value=lead.get("role", ""))
        ws.cell(row=row, column=5, value=lead.get("location", ""))
        fit_cell = ws.cell(row=row, column=6, value=lead.get("fit", ""))
        fit_cell.alignment = CENTER
        fit_cell.font = Font(name=FONT_NAME, bold=True, size=11)
        band_cell = ws.cell(row=row, column=7, value=band)
        band_cell.fill = BAND_FILLS.get(band, BAND_FILLS["Weak"])
        band_cell.alignment = CENTER
        band_cell.font = Font(name=FONT_NAME, bold=True, size=10)
        ws.cell(row=row, column=8, value=lead.get("top_signal", ""))
        ws.cell(row=row, column=9, value=lead.get("rationale", ""))
        email_val = lead.get("email", "")
        if lead.get("email_status"):
            email_val = f"{email_val} ({lead['email_status']})" if email_val else ""
        ws.cell(row=row, column=10, value=email_val)
        ws.cell(row=row, column=11, value=lead.get("next_step", ""))
        ws.cell(row=row, column=12, value=lead.get("linkedin_url", ""))

        for col in range(1, 13):
            c = ws.cell(row=row, column=col)
            if col not in (1, 6, 7):
                c.alignment = WRAP
                if not c.font.bold:
                    c.font = BODY_FONT
            c.border = BORDER
        ws.row_dimensions[row].height = 72

    ws.freeze_panes = "A2"

    # ---------------- Per-lead detail tabs ----------------
    used_tabs = {"overview"}
    for lead in leads_sorted:
        name = lead.get("name", "lead")
        band = band_for(lead.get("fit"), lead.get("band"))
        tab = wb.create_sheet(sanitize_tab(name, used_tabs))
        tab.column_dimensions["A"].width = 22
        tab.column_dimensions["B"].width = 100

        tab["A1"] = name
        tab["A1"].font = Font(name=FONT_NAME, bold=True, size=16, color="1F3A5F")
        tab.merge_cells("A1:B1")
        subtitle = f"{lead.get('role','')} at {lead.get('company','')}  ·  {lead.get('location','')}"
        tab["A2"] = subtitle.strip(" ·")
        tab["A2"].font = Font(name=FONT_NAME, size=11, italic=True, color="555555")
        tab.merge_cells("A2:B2")

        def put(row, label, value, fill=None, bold_body=False):
            lc = tab.cell(row=row, column=1, value=label)
            lc.font = SECTION_FONT
            lc.fill = SECTION_FILL
            lc.alignment = WRAP
            cell = tab.cell(row=row, column=2, value=value)
            cell.alignment = WRAP
            cell.font = Font(name=FONT_NAME, bold=True, size=11) if bold_body else BODY_FONT
            if fill is not None:
                cell.fill = fill
            return cell

        row = 3
        put(row, "Fit (0-5)", lead.get("fit", ""), bold_body=True)
        row += 1
        put(row, "Band", band, BAND_FILLS.get(band), bold_body=True)
        row += 1
        put(row, "Top signal", lead.get("top_signal", ""))
        tab.row_dimensions[row].height = row_height(lead.get("top_signal"))
        row += 1
        put(row, "Rationale", lead.get("rationale", ""))
        tab.row_dimensions[row].height = row_height(lead.get("rationale"))
        row += 1
        put(row, "Next step", lead.get("next_step", ""))
        tab.row_dimensions[row].height = row_height(lead.get("next_step"))
        row += 1

        email_val = lead.get("email", "")
        if lead.get("email_status"):
            email_val = f"{email_val} ({lead['email_status']})" if email_val else lead["email_status"]
        put(row, "Email", email_val)
        row += 1
        put(row, "LinkedIn", lead.get("linkedin_url", ""))
        row += 1
        if lead.get("match_note"):
            put(row, "Match note", lead.get("match_note"))
            row += 1

        # Research dossier
        row += 1
        hc = tab.cell(row=row, column=1, value="RESEARCH")
        hc.font = Font(name=FONT_NAME, bold=True, size=12, color="FFFFFF")
        hc.fill = HEADER_FILL
        tab.cell(row=row, column=2).fill = HEADER_FILL
        tab.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
        row += 1
        research = lead.get("research", "")
        put(row, "Dossier", research)
        tab.row_dimensions[row].height = row_height(research, per_chars=8, hi=400)
        row += 1

        sources = lead.get("sources", []) or []
        if sources:
            put(row, "Sources", "\n".join(str(s) for s in sources))
            tab.row_dimensions[row].height = row_height("\n".join(map(str, sources)), per_chars=10)
            row += 1

        tab.freeze_panes = "A3"

    slug = meta.get("project_slug") or f"{_date.today().isoformat()}-icp-lead-ranking"
    out_path = os.path.join(out_dir, f"{slug}-leads-ranked.xlsx")
    wb.save(out_path)
    print(out_path)


if __name__ == "__main__":
    main()
