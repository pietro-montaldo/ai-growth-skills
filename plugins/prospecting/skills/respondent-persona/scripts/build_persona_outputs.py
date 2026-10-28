#!/usr/bin/env python3
"""Build the respondent-persona spread outputs from respondents.json.

Needs openpyxl (pip install openpyxl).

Usage:
    python3 build_persona_outputs.py respondents.json output_dir [--anonymous]

--anonymous leaves the name and email columns out of the Respondents sheet,
for a copy you can share. Without it, the Respondents sheet keeps them so you
can audit the read; keep that file private.

Writes:
    <slug>-spread.xlsx   Spread sheet (counts + % per bucket per dimension) + Respondents sheet
    <slug>-spread.csv    The spread, flat

The slug comes from meta.project_slug; falls back to the json filename stem.
"""

import json
import sys
import csv
from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

# Dimension -> ordered bucket list. "unknown" is appended to every dimension.
DIMENSIONS = {
    "seniority": ["executive", "senior", "mid", "junior", "student"],
    "function": ["marketing", "sales", "product", "engineering", "data",
                 "operations", "finance", "hr", "consulting", "founder_exec",
                 "creative", "other"],
    "industry": ["software_saas", "agency_services", "ecommerce_retail",
                 "finance_fintech", "healthcare", "education", "media_content",
                 "manufacturing", "nonprofit_gov", "other"],
    "company_size": ["solo", "micro", "small", "mid", "large", "enterprise"],
    "region": ["north_america", "western_europe", "eastern_europe", "latam",
               "mena", "sub_saharan_africa", "south_asia", "east_asia",
               "oceania"],
}

RESP_COLS = ["name", "email", "company", "company_source", "seniority",
             "function", "industry", "company_size", "region", "confidence",
             "evidence"]

HEADER_FILL = PatternFill("solid", fgColor="1A1A1A")
HEADER_FONT = Font(color="FFFFFF", bold=True)
DIM_FILL = PatternFill("solid", fgColor="EDEDEB")
DIM_FONT = Font(bold=True)


def pct(n, total):
    return round(100.0 * n / total, 1) if total else 0.0


def compute_spread(respondents):
    """Return list of rows: (dimension, bucket, count, pct) plus the modal bucket per dimension."""
    total = len(respondents)
    rows = []
    modes = {}
    for dim, buckets in DIMENSIONS.items():
        counts = Counter(r.get(dim, "unknown") or "unknown" for r in respondents)
        ordered = buckets + ["unknown"]
        # include any unexpected value the data carried, after the known buckets
        for v in counts:
            if v not in ordered:
                ordered.insert(-1, v)
        # modal bucket excluding unknown, for the persona convenience
        non_unknown = {b: counts.get(b, 0) for b in ordered if b != "unknown"}
        mode = max(non_unknown, key=non_unknown.get) if non_unknown and max(non_unknown.values()) > 0 else "unknown"
        modes[dim] = (mode, counts.get(mode, 0), pct(counts.get(mode, 0), total),
                      pct(counts.get("unknown", 0), total))
        for b in ordered:
            c = counts.get(b, 0)
            if c == 0 and b not in buckets and b != "unknown":
                continue
            rows.append((dim, b, c, pct(c, total)))
    return rows, modes, total


def write_xlsx(respondents, spread_rows, modes, total, out_path, anonymous=False):
    wb = Workbook()

    # --- Spread sheet ---
    ws = wb.active
    ws.title = "Spread"
    ws.append(["Dimension", "Bucket", "Count", "% of researched"])
    for c in ws[1]:
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
    last_dim = None
    for dim, bucket, count, p in spread_rows:
        show_dim = "" if dim == last_dim else dim
        ws.append([show_dim, bucket, count, p])
        if show_dim:
            ws.cell(ws.max_row, 1).font = DIM_FONT
        last_dim = dim
    ws.append([])
    ws.append([f"Researched respondents: {total}"])
    ws.cell(ws.max_row, 1).font = DIM_FONT
    for col, w in zip("ABCD", [16, 20, 10, 16]):
        ws.column_dimensions[col].width = w

    # --- Average-respondent summary (top of a small sheet) ---
    ws2 = wb.create_sheet("Average respondent")
    ws2.append(["Dimension", "Modal bucket", "Count", "% of total", "% unknown"])
    for c in ws2[1]:
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
    for dim in DIMENSIONS:
        mode, cnt, p, unk = modes[dim]
        ws2.append([dim, mode, cnt, p, unk])
    for col, w in zip("ABCDE", [16, 18, 10, 12, 12]):
        ws2.column_dimensions[col].width = w

    # --- Respondents sheet ---
    ws3 = wb.create_sheet("Respondents")
    cols = [c for c in RESP_COLS if not (anonymous and c in ("name", "email"))]
    ws3.append(cols)
    for c in ws3[1]:
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
    for r in respondents:
        ws3.append([r.get(k, "") for k in cols])
    width_map = dict(zip(RESP_COLS, [20, 26, 22, 14, 12, 14, 16, 14, 16, 11, 40]))
    widths = [width_map[c] for c in cols]
    for col, w in zip("ABCDEFGHIJK", widths):
        ws3.column_dimensions[col].width = w
    ws3.freeze_panes = "A2"

    wb.save(out_path)


def write_csv(spread_rows, total, out_path):
    with open(out_path, "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["dimension", "bucket", "count", "pct_of_researched"])
        for row in spread_rows:
            w.writerow(row)
        w.writerow([])
        w.writerow(["researched_total", total])


def main():
    args = [a for a in sys.argv[1:] if a != "--anonymous"]
    anonymous = "--anonymous" in sys.argv[1:]
    if len(args) != 2:
        print(__doc__)
        sys.exit(1)
    json_path = Path(args[0])
    out_dir = Path(args[1])
    out_dir.mkdir(parents=True, exist_ok=True)

    data = json.loads(json_path.read_text())
    respondents = data.get("respondents", [])
    if not respondents:
        print("No respondents in JSON.")
        sys.exit(1)
    slug = data.get("meta", {}).get("project_slug") or json_path.stem

    spread_rows, modes, total = compute_spread(respondents)

    xlsx_path = out_dir / f"{slug}-spread.xlsx"
    csv_path = out_dir / f"{slug}-spread.csv"
    write_xlsx(respondents, spread_rows, modes, total, xlsx_path, anonymous)
    write_csv(spread_rows, total, csv_path)

    print(f"Wrote {xlsx_path}")
    print(f"Wrote {csv_path}")
    print(f"Researched respondents: {total}")
    print("Modal buckets:")
    for dim in DIMENSIONS:
        mode, cnt, p, unk = modes[dim]
        print(f"  {dim:14s} {mode:16s} {p:5.1f}%  (unknown {unk:.1f}%)")


if __name__ == "__main__":
    main()
