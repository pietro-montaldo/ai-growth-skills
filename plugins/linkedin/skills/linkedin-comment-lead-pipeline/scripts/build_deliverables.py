#!/usr/bin/env python3
"""Build the ranked lead list (.csv, plus .xlsx if openpyxl is installed) and the
fit 4-5 DM drafts (.md) from a leads.json file.

Usage:
    python3 build_deliverables.py <leads.json> <output_dir>

leads.json schema:
{
  "meta": {"post_urls": [], "post_topic": "", "offer": "", "offer_url": "", "slug": ""},
  "leads": [
    {"name", "type", "headline", "profile_url", "engagement", "posts_engaged", "comment",
     "enriched_via", "enriched_name", "title", "company", "location", "email",
     "email_status", "match_note", "fit", "rationale", "dm"}
  ]
}

The CSV needs only the standard library. The .xlsx is a nicer view of the same data
and is skipped (with a note) when openpyxl is not installed.
"""
import csv
import json
import os
import sys

COLS = ["Rank", "Fit (0-5)", "Name", "Type", "Title", "Company", "Location",
        "Email", "Email Status", "Why this score", "Engagement", "Posts engaged",
        "LinkedIn URL", "Comment", "Match Note"]


def g(d, k, default=""):
    v = d.get(k, default)
    return default if v is None else v


def as_int(v):
    try:
        return int(v or 0)
    except (TypeError, ValueError):
        return 0


def row(i, r):
    return [i, r["_fit"], g(r, "name"), g(r, "type"), g(r, "title"), g(r, "company"),
            g(r, "location"), g(r, "email"), g(r, "email_status"), g(r, "rationale"),
            g(r, "engagement"), g(r, "posts_engaged"), g(r, "profile_url"),
            g(r, "comment"), g(r, "match_note")]


def write_csv(base, leads):
    path = base + "-leads-ranked.csv"
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(COLS)
        for i, r in enumerate(leads, 1):
            w.writerow(row(i, r))
    return path


def write_xlsx(base, leads, meta):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    except ImportError:
        return None

    wb = Workbook()
    ws = wb.active
    ws.title = "ICP Ranked"
    head = Font(name="Arial", bold=True, color="FFFFFF", size=11)
    hfill = PatternFill("solid", start_color="2E5B8A")
    body = Font(name="Arial", size=10)
    top = Alignment(vertical="top", wrap_text=True)
    ctr = Alignment(vertical="top", horizontal="center", wrap_text=True)
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    fit_colors = {5: "C6EFCE", 4: "D9EAD3", 3: "FFF2CC", 2: "FCE4D6", 1: "F4CCCC", 0: "EDEDED"}

    ws["A1"] = "LinkedIn engagers ranked by ICP fit"
    ws["A1"].font = Font(name="Arial", bold=True, size=13)
    posts = ", ".join(meta.get("post_urls") or [])
    ws["A2"] = f'Topic: {g(meta, "post_topic", "(post)")}  |  {posts}'
    ws["A2"].font = Font(name="Arial", size=9, color="808080")

    hr = 4
    centered = ("Rank", "Fit (0-5)", "Email Status", "Posts engaged")
    for c, h in enumerate(COLS, 1):
        cell = ws.cell(row=hr, column=c, value=h)
        cell.font = head
        cell.fill = hfill
        cell.border = border
        cell.alignment = Alignment(vertical="center", wrap_text=True,
                                   horizontal="center" if h in centered else "left")

    for i, r in enumerate(leads):
        rr = hr + 1 + i
        for c, v in enumerate(row(i + 1, r), 1):
            cell = ws.cell(row=rr, column=c, value=v)
            cell.font = body
            cell.border = border
            cell.alignment = ctr if COLS[c - 1] in centered else top
        fit_cell = ws.cell(row=rr, column=2)
        fit_cell.fill = PatternFill("solid", start_color=fit_colors.get(r["_fit"], "EDEDED"))
        fit_cell.font = Font(name="Arial", size=11, bold=True)
        if g(r, "email"):
            ws.cell(row=rr, column=8).fill = PatternFill("solid", start_color="E2EFDA")
        ws.cell(row=rr, column=13).font = Font(name="Arial", size=10, color="0563C1")

    widths = [5, 7, 22, 10, 34, 24, 24, 30, 11, 48, 16, 9, 38, 50, 32]
    for idx, wd in enumerate(widths):
        ws.column_dimensions[chr(ord("A") + idx)].width = wd
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A{hr}:{chr(ord('A') + len(COLS) - 1)}{hr + len(leads)}"
    path = base + "-leads-ranked.xlsx"
    wb.save(path)
    return path


def write_dms(base, leads, meta):
    dms = [r for r in leads if r["_fit"] >= 4 and g(r, "dm").strip()]
    if not dms:
        return None, 0
    offer = g(meta, "offer") or g(meta, "offer_url")
    lines = ["# DM drafts for fit 4 and 5", "",
             f"**Offer:** {offer}" + (f' ({g(meta, "offer_url")})' if g(meta, "offer") and g(meta, "offer_url") else ""),
             f'**Source posts:** {", ".join(meta.get("post_urls") or [])}',
             "**Note:** drafts only. Open each profile and send them yourself. Nothing was sent.", ""]
    for fit in (5, 4):
        grp = [r for r in dms if r["_fit"] == fit]
        if not grp:
            continue
        lines += [f"## Fit {fit}", ""]
        for r in grp:
            tag = g(r, "title") or g(r, "headline")
            lines.append(f'**{g(r, "name")}**' + (f" | {tag}" if tag else "") + "  ")
            lines.append(f'LinkedIn: {g(r, "profile_url") or "(no URL)"}')
            lines.append("")
            lines.append(g(r, "dm").strip())
            lines.append("")
    path = base + "-dms.md"
    with open(path, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))
    return path, len(dms)


def main():
    if len(sys.argv) < 3:
        print("Usage: python3 build_deliverables.py <leads.json> <output_dir>")
        sys.exit(1)
    with open(sys.argv[1], encoding="utf-8") as f:
        data = json.load(f)
    outdir = sys.argv[2]
    meta = data.get("meta", {})
    leads = data.get("leads", [])
    os.makedirs(outdir, exist_ok=True)
    base = os.path.join(outdir, g(meta, "slug", "linkedin-engager-leads"))

    for r in leads:
        r["_fit"] = as_int(g(r, "fit", 0))
        r["_mail"] = 1 if g(r, "email") else 0
        r["_comment"] = 1 if "comment" in str(g(r, "engagement")).lower() else 0
        r["_posts"] = as_int(g(r, "posts_engaged", 0))
    # Best fit first, then leads with an email, then commenters, then most posts engaged.
    leads.sort(key=lambda x: (-x["_fit"], -x["_mail"], -x["_comment"], -x["_posts"],
                              str(g(x, "name")).lower()))

    missing_url = [g(r, "name") for r in leads if r["_fit"] >= 4 and g(r, "dm").strip()
                   and not g(r, "profile_url")]

    csv_path = write_csv(base, leads)
    xlsx_path = write_xlsx(base, leads, meta)
    dm_path, n_dms = write_dms(base, leads, meta)

    n5 = sum(1 for r in leads if r["_fit"] == 5)
    n4 = sum(1 for r in leads if r["_fit"] == 4)
    print(f"leads: {len(leads)} | fit 5: {n5} | fit 4: {n4} | DMs: {n_dms}")
    print("saved:", csv_path)
    if xlsx_path:
        print("saved:", xlsx_path)
    else:
        print("note: openpyxl is not installed, so no .xlsx was written. The .csv has the same data.")
    if dm_path:
        print("saved:", dm_path)
    if missing_url:
        print("warning: these DM leads have no LinkedIn URL:", ", ".join(missing_url))


if __name__ == "__main__":
    main()
